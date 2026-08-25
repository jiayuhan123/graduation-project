from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path(r"E:\graduation-project\primary-math-platform-feature-plan-.docx")
TARGET = Path(r"E:\graduation-project\primary-math-platform-feature-plan-updated.docx")


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace visible text while preserving the first run's formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text}")


def clone_paragraph_before(template: Paragraph, marker: Paragraph, text: str) -> Paragraph:
    new_element = deepcopy(template._p)
    marker._p.addprevious(new_element)
    paragraph = Paragraph(new_element, marker._p.getparent())
    set_paragraph_text(paragraph, text)
    return paragraph


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def append_cloned_row(table):
    new_row = deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_row)
    return table.rows[-1]


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    shutil.copy2(SOURCE, TARGET)
    document = Document(TARGET)

    # Preserve the existing visual system and update the opening positioning.
    set_paragraph_text(
        find_paragraph(document, "小学计算智能学习开放平台"),
        "小学数学智能学习与游戏化对战平台",
    )
    set_cell_text(
        document.tables[0].cell(0, 0),
        "系统定位\n面向不同学校及独立学生，建设集练习、拍照批改、错因诊断、掌握度更新、个性化训练、游戏化激励、实时对战和开放接口于一体的通用小学数学智能学习平台。学生无需加入学校或班级即可使用核心学习与公共对战功能；加入组织后可进一步使用教师作业、班级统计和校内竞赛。",
    )

    replacements = {
        "家长拍摄整页口算作业后，系统自动切分题目、识别答案并标出错题；数学结果由程序重新计算和校验。当某个数字或符号识别不清时，系统先提示家长或教师确认，确认前不把该结果写入学生学情。":
            "学生拍摄整页数学作业后，系统自动切分题目、识别答案并标出错题；数学结果由程序重新计算和校验。当某个数字或符号识别不清时，系统先提示学生确认，必要时由教师复核，确认前不把该结果写入学生学情。",
        "孩子每完成一道题，系统都会更新对应知识点的掌握情况。家长和教师可以看到哪些内容已经熟练、哪些还不稳定、哪些虽然会做但速度偏慢，以及哪些内容长时间未练、可能出现遗忘。":
            "学生每完成一道题，系统都会更新对应知识点的掌握情况。学生本人和有管理权限的教师可以看到哪些内容已经熟练、哪些还不稳定、哪些虽然会做但速度偏慢，以及哪些内容长时间未练、可能出现遗忘。",
        "程序负责答案计算、规则校验和结果确认；AI主要负责儿童化讲解、换一种方式说明、生成家长学习报告、整理班级共性问题以及形成教师作业草案。这种分工可以使用AI的表达能力，同时降低AI计算错误对学习结果的影响。":
            "程序负责答案计算、规则校验和结果确认；AI主要负责儿童化讲解、换一种方式说明、生成学生学习报告、整理班级共性问题以及形成教师作业草案。这种分工可以使用AI的表达能力，同时降低AI计算错误对学习结果的影响。",
        "学生端不提供完全开放的AI聊天，减少直接代答、代写风险；学生姓名、班级、作业图片和学习记录按照权限隔离。教师生成作业、发布训练计划等重要操作必须经过人工确认，并保留操作记录。":
            "学生端不提供完全开放的AI聊天，减少直接代答、代写风险；学生身份、学校与班级归属、作业图片和学习记录按照权限与数据范围隔离。教师生成作业、发布训练计划等重要操作必须经过人工确认，并保留操作记录。",
    }
    for old_text, new_text in replacements.items():
        set_paragraph_text(find_paragraph(document, old_text), new_text)

    # Reuse the established feature-heading/body formatting.
    heading_template = find_paragraph(document, "1. 可信的整页拍照批改")
    body_template = find_paragraph(
        document,
        "学生拍摄整页数学作业后，系统自动切分题目、识别答案并标出错题；数学结果由程序重新计算和校验。当某个数字或符号识别不清时，系统先提示学生确认，必要时由教师复核，确认前不把该结果写入学生学情。",
    )
    section_template = find_paragraph(document, "一、核心特色（重点）")
    module_heading = find_paragraph(document, "二、主要功能模块（简版）")

    clone_paragraph_before(
        heading_template,
        module_heading,
        "9. 独立学生可直接使用，多学校通用",
    )
    clone_paragraph_before(
        body_template,
        module_heading,
        "学生注册后即可选择年级、学期和教材版本，独立使用练习、批改、错因分析、个性化推荐、游戏激励和公共对战功能。学校与班级是可选组织归属；加入组织后，再获得教师作业、班级统计和校内竞赛等协同能力。",
    )
    clone_paragraph_before(
        heading_template,
        module_heading,
        "10. 将学习过程游戏化，并支持一对一和组队PK",
    )
    clone_paragraph_before(
        body_template,
        module_heading,
        "积分、经验、等级、徽章、连续学习和排行榜用于激励持续学习；一对一PK支持随机匹配、邀请和房间对战，组队PK支持建队、成员管理、队伍匹配与团体结算。对战以正确率为主、用时为辅，并保存完整答题和比赛记录。",
    )

    # Rewrite the concise module table without introducing additional human roles.
    table = document.tables[1]
    rows = [
        (
            "功能端/模块",
            "主要功能",
        ),
        (
            "学生端",
            "学生登录与个人信息；独立使用或申请加入学校、班级；选择年级、学期和教材版本；每日练习、专项练习、限时练习和教师作业；在线答题、计时、自动判题与整页拍照批改；查看错因讲解、相似题、错题本、学习记录、个性化推荐、掌握度和学习报告；使用积分、等级、徽章、任务和排行榜；参与一对一及组队PK。",
        ),
        (
            "教师管理端",
            "教师登录、班级和学生管理；创建并布置班级作业、个人作业和分层作业；查看完成情况、答题详情、正确率、高频错题和薄弱知识点；复核识别不清或无法确定的答案；AI辅助生成作业草案与训练计划，由教师确认后发布；组织班级PK和队伍竞赛；导出学习报告。",
        ),
        (
            "管理员端",
            "统一管理用户、角色、学校、年级和班级，并通过数据范围区分平台管理与学校管理；维护教材版本、知识点、题型、题目、难度、出题规则和错因分类；配置推荐、AI与OCR服务；配置积分、等级、徽章、排行榜、匹配和计分规则；管理第三方应用、接口权限、内容安全、审计日志、监控和异常告警。",
        ),
        (
            "游戏化与PK模块",
            "积分、经验、等级、徽章、成就、连续学习、每日任务和排行榜；一对一随机匹配、邀请匹配和房间匹配；队伍创建、加入、退出和成员管理；组队匹配、班级对抗、对局题目、实时答题、计分、结算、个人贡献和历史战绩。",
        ),
        (
            "开放平台能力",
            "第三方应用不是新的用户角色，而是由管理员创建和审核的机器访问主体；提供应用、密钥、接口权限、调用额度、限流、调用日志和错误追踪，并开放出题、判题、图片批改、错因分析、掌握度查询、个性化推荐和AI讲解接口。",
        ),
    ]
    while len(table.rows) < len(rows):
        append_cloned_row(table)
    for row_index, (left, right) in enumerate(rows):
        set_cell_text(table.cell(row_index, 0), left)
        set_cell_text(table.cell(row_index, 1), right)

    # Add detailed scope decisions before the original project-focus section.
    project_heading = find_paragraph(document, "三、项目重点")

    clone_paragraph_before(section_template, project_heading, "三、角色与使用边界")
    role_sections = [
        (
            "1. 学生",
            "学生可以独立注册和使用平台，也可以申请加入学校与班级。核心功能包括在线练习、拍照批改、错因讲解、错题本、知识点掌握度、个性化推荐、学习报告、游戏化激励以及一对一和组队PK。",
        ),
        (
            "2. 教师",
            "教师负责自己数据范围内的班级、学生、题目和作业，查看学习统计与薄弱知识点，复核不确定结果，发布AI辅助生成的训练计划，并组织班级或队伍竞赛。",
        ),
        (
            "3. 管理员",
            "管理员负责平台或学校范围内的组织、用户、课程资源、题库、智能服务、游戏规则、PK规则、开放接口、安全审计与系统运行管理。平台管理员和学校管理员使用同一角色，通过数据权限范围进行区分。",
        ),
        (
            "4. 角色范围原则",
            "系统的人类用户角色暂定为学生、教师和管理员三类。家长不作为独立登录角色；第三方应用作为机器访问主体，由管理员管理，不增加开发者用户角色。",
        ),
    ]
    for title, body in role_sections:
        clone_paragraph_before(heading_template, project_heading, title)
        clone_paragraph_before(body_template, project_heading, body)

    clone_paragraph_before(section_template, project_heading, "四、游戏化与PK功能")
    game_sections = [
        (
            "1. 游戏化激励",
            "将练习完成、错题订正、知识点掌握和持续学习转化为积分、经验、等级、徽章、成就和排行榜。奖励以学习进步和正确性为主，避免单纯鼓励速度。",
        ),
        (
            "2. 一对一PK",
            "支持随机匹配、同学邀请和房间匹配，并按照年级、知识点与能力范围匹配对手。双方使用相同题目或同难度等价题目，按正确率优先、答题用时辅助的规则完成实时计分和结果结算。",
        ),
        (
            "3. 组队PK",
            "支持队伍创建、加入、退出、成员管理和队伍匹配，可开展班级对抗或自定义房间竞赛，并按照总分、平均分或接力答题规则计算队伍成绩与个人贡献。",
        ),
        (
            "4. 未成年人安全与竞赛公平",
            "学生端不提供自由聊天，邀请、组队和对局受到系统规则与教师管理约束。对局保存参赛者、题目、答案、得分、异常状态和结算结果，为申诉处理与公平性检查提供依据。",
        ),
    ]
    for title, body in game_sections:
        clone_paragraph_before(heading_template, project_heading, title)
        clone_paragraph_before(body_template, project_heading, body)

    clone_paragraph_before(section_template, project_heading, "五、通用性与组织关系")
    general_sections = [
        (
            "1. 多学校通用",
            "平台支持多个学校，各学校共享统一业务能力，但学校、班级、作业和学习数据按照数据权限隔离。学校差异通过年级、学期、教材版本、课程资源和管理配置体现，不针对单个学校编写独立业务分支。",
        ),
        (
            "2. 学校和班级为可选归属",
            "学生档案独立存在，不以加入学校或班级作为使用前提。学生与学校、学生与班级通过可选关联记录建立关系；未加入组织时仍可学习和参与公共PK，关联记录同时保存入校、升班、转学和离校历史。",
        ),
        (
            "3. 年级差异通过内容适用范围表达",
            "学生独立使用时自行选择年级、学期和教材版本；加入班级后可由班级配置覆盖默认选择。题目通过年级、学期、教材版本、知识点和难度确定适用范围，而不是通过学校专属代码区分。",
        ),
        (
            "4. 开放能力不改变角色模型",
            "平台继续保留第三方系统接入能力。第三方应用使用应用标识与密钥调用受授权接口，由管理员负责审核、授权、限流和审计，因此不需要新增开发者角色。",
        ),
    ]
    for title, body in general_sections:
        clone_paragraph_before(heading_template, project_heading, title)
        clone_paragraph_before(body_template, project_heading, body)

    set_paragraph_text(project_heading, "六、项目重点")
    set_paragraph_text(
        find_paragraph(
            document,
            "本项目的重点不是单独证明AI能够识别一道数学题，而是建设一套可落地的学习闭环：批改结果经过确定性程序校验，具体错因进入学生长期学情，学情继续影响下一轮训练，教师和家长能够看懂变化，第三方应用也能够通过标准接口复用这些能力。",
        ),
        "本项目的重点不是单独证明AI能够识别一道数学题，而是建设一套可落地、可持续使用的学习平台：批改结果经过确定性程序校验，具体错因进入学生长期学情，学情继续影响下一轮训练；游戏化激励和PK提高学习参与度；独立学生与在校学生共享核心能力，多学校通过统一模型和数据权限获得通用支持；第三方应用可通过标准接口复用学习闭环。实施时以学习业务闭环、三类角色和对战功能为主，开放平台与复杂运营能力在核心功能完成后继续完善。",
    )
    set_cell_text(
        document.tables[2].cell(0, 0),
        "最终形成\n学习闭环：做题 → 可信批改 → 找到具体错因 → 更新掌握情况 → 安排针对性练习 → 再次验证\n激励闭环：完成学习 → 获得积分与成长 → 解锁成就 → 保持持续学习\n对战闭环：匹配或组队 → 公平答题 → 实时计分 → 结果结算 → 战绩与奖励沉淀",
    )

    properties = document.core_properties
    properties.title = "小学数学智能学习与游戏化对战平台功能规划"
    properties.subject = "三类角色、独立学生、多学校通用、游戏化与PK、开放接口"
    properties.comments = "依据毕业设计沟通结论更新，原文件已保留。"

    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
