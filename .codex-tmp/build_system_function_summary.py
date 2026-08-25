from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\graduation-project")
OUTPUT = ROOT / "文档" / "系统功能总结.docx"

DESIGN_PRESET = "compact_reference_guide"
HEADER_PATTERN = "editorial_cover"
EAST_ASIA_FONT_OVERRIDE = "Microsoft YaHei"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
BODY = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "B8C6D1"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, *, size=None, bold=None, italic=None, color=None, east_asia=EAST_ASIA_FONT_OVERRIDE):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def configure_style(style, *, size, color=BODY, bold=False, before=0, after=0, line=1.0, align=None):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = rgb(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT_OVERRIDE)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field_run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    field_run._r.append(fld_char_begin)
    field_run._r.append(instr_text)
    field_run._r.append(fld_char_end)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, *, ordered: bool):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if not ordered:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
        r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, *, bold_lead=None):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_el)
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, size=11, bold=True, color=BODY)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11, color=BODY)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=BODY)
    return p


def add_body(doc, text, *, bold_lead=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style="Normal")
    p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, size=11, bold=True, color=BODY)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11, color=BODY)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=BODY)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}[level],
            bold=True,
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level],
        )
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, color="D5DEE8", size=4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    lead = p.add_run(title)
    set_run_font(lead, size=11, bold=True, color=NAVY)
    detail = p.add_run(text)
    set_run_font(detail, size=11, color=BODY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_module_overview(doc):
    modules = [
        ("1", "学生与学习档案", "学生独立使用、可选加入学校或班级，维护年级、学期、教材版本和学习记录。"),
        ("2", "教材与知识体系", "管理教材版本、教材章节和知识点映射，为跨学校、跨教材的内容适配提供统一基础。"),
        ("3", "题库与出题", "维护题型、难度、标准答案、知识点和适用范围，支持规则出题与个性化组题。"),
        ("4", "练习与作答", "支持个性化口算、专项练习、限时练习和在线作答，完整记录题目、答案、用时与得分。"),
        ("5", "拍照辅助批改", "识别题目与答案，完成可确认的批改；对低可信结果提示学生确认或进入教师复核。"),
        ("6", "错因与个性化推荐", "归纳错误类型、更新知识点掌握度，生成针对性练习并通过再次验证确认学习效果。"),
        ("7", "教师作业与报告", "教师管理班级和学生，布置作业、查看完成情况、分析薄弱知识点并形成学习报告。"),
        ("8", "游戏闯关与激励", "通过单人闯关、积分、经验、等级、徽章、成就、每日任务和排行榜促进持续练习。"),
        ("9", "一对一与组队对战", "支持匹配、房间、实时答题、计分、结算、队伍管理和个人贡献记录。"),
        ("10", "平台运营管理", "管理学校、班级、用户、教材内容、题库、业务规则、异常处理和运营统计。"),
        ("11", "开放平台", "面向第三方应用开放出题、批改、识别、诊断、掌握度查询、推荐和讲解等业务能力。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ("序号", "业务模块", "第一版主要能力")
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        set_cell_shading(cell, LIGHT_BLUE)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, size=10.5, bold=True, color=NAVY)
    repeat_table_header(table.rows[0])

    for number, name, description in modules:
        cells = table.add_row().cells
        values = (number, name, description)
        for i, value in enumerate(values):
            cells[i].text = value
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, size=10, bold=(i == 1), color=BODY)
    for row in table.rows:
        prevent_row_split(row)
    set_table_geometry(table, [720, 2280, 6360], indent_dxa=120)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_style(doc.styles["Normal"], size=11, color=BODY, after=6, line=1.25)
    configure_style(doc.styles["Title"], size=28, color=NAVY, bold=True, before=0, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    configure_style(doc.styles["Subtitle"], size=13, color=DARK_BLUE, before=0, after=10, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    configure_style(doc.styles["Heading 1"], size=16, color=BLUE, bold=True, before=18, after=10, line=1.0)
    configure_style(doc.styles["Heading 2"], size=13, color=BLUE, bold=True, before=14, after=7, line=1.0)
    configure_style(doc.styles["Heading 3"], size=12, color=DARK_BLUE, bold=True, before=10, after=5, line=1.0)
    configure_style(doc.styles["List Paragraph"], size=11, color=BODY, after=4, line=1.25)

    title_p_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style_name].paragraph_format.keep_with_next = True
        doc.styles[style_name].paragraph_format.keep_together = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("小学数学智能学习与游戏对战平台")
    set_run_font(left, size=8.5, color=MUTED)
    right = p.add_run("\t第一版业务功能说明")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)


def build_document():
    doc = Document()
    configure_document(doc)
    bullet_id = add_numbering_definition(doc, ordered=False)
    decimal_id = add_numbering_definition(doc, ordered=True)

    properties = doc.core_properties
    properties.title = "小学数学智能学习与游戏对战平台 - 系统功能总结"
    properties.subject = "第一版业务功能范围"
    properties.author = "毕业设计项目组"
    properties.keywords = "小学数学, 个性化学习, 游戏闯关, 对战, 开放平台"

    # Editorial-cover opening block.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("毕业设计 · 第一版业务范围基线")
    set_run_font(run, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph("小学数学智能学习与游戏对战平台", style="Title")
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    subtitle = doc.add_paragraph("系统功能总结", style="Subtitle")
    subtitle.paragraph_format.keep_with_next = True
    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.space_after = Pt(22)
    run = desc.add_run("通用型小学数学学习、游戏闯关、协作对战与开放能力平台")
    set_run_font(run, size=12, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(32)
    run = meta.add_run("版本 V1.0  |  业务功能说明  |  2026年8月")
    set_run_font(run, size=10.5, color=MUTED)

    add_callout(
        doc,
        "第一版定位：",
        "面向小学生、教师和管理员，围绕“练习—批改—诊断—推荐—验证”建立可持续学习闭环，同时提供游戏闯关、一对一与组队对战，并通过开放平台向第三方应用输出核心学习能力。",
    )
    doc.add_page_break()

    add_heading(doc, "1 第一版定位与范围", 1)
    add_body(
        doc,
        "本系统是一套面向小学数学学习场景的通用业务平台，不针对某一所学校定制。学生可以不加入学校或班级，直接选择年级、学期和教材版本使用核心学习功能；加入组织后，可进一步接收教师作业、参与班级统计与校内活动。",
    )
    add_body(
        doc,
        "第一版重点不是单独验证某一种识别或智能能力，而是把内容组织、练习作答、辅助批改、错因诊断、掌握度更新、个性化推荐、再次验证、教师教学、游戏激励和开放能力连接成完整业务流程。",
    )
    add_callout(
        doc,
        "核心约束：",
        "第一版业务角色仅包括学生、教师和管理员；家长不作为独立使用角色；第三方开发者通过开放平台的应用主体接入，不增加新的人工业务角色。",
    )

    add_heading(doc, "2 业务角色与组织使用方式", 1)
    add_heading(doc, "2.1 学生", 2)
    for text in [
        "注册后即可独立使用核心学习、练习、拍照辅助批改、错题、掌握度、推荐、闯关和公共对战功能。",
        "可选加入学校和班级；加入后继承班级的默认年级、学期和教材设置，并接收教师布置的作业与训练计划。",
        "查看个人学习记录、错题、薄弱知识点、阶段报告、积分、等级、徽章、排行榜和对战记录。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "2.2 教师", 2)
    for text in [
        "管理本人负责范围内的班级与学生，查看学习进度、作答情况和薄弱知识点。",
        "从题库选题或按知识点、难度和题量生成作业，发布普通作业、分层作业和训练计划。",
        "复核低可信批改结果，查看作业统计与学习报告，并组织班级闯关或对战活动。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "2.3 管理员", 2)
    for text in [
        "维护学校、年级、班级和用户等基础业务数据。",
        "维护教材版本、教材章节、知识点、题型、题目、答案、难度和业务规则。",
        "配置游戏、闯关、对战、推荐、识别和辅助批改等运营参数，处理异常业务记录。",
        "管理第三方应用接入、开放能力范围、调用额度和业务使用情况。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "3 核心学习闭环", 1)
    add_callout(
        doc,
        "学习闭环：",
        "选择或生成练习 → 在线作答或拍照提交 → 辅助批改与结果确认 → 形成错因记录 → 更新知识点掌握度 → 推荐针对性练习 → 再次验证。",
    )
    for text in [
        "练习来源包括学生自主练习、个性化口算、教师作业、闯关题目、对战题目和拍照任务。",
        "每次作答记录学生答案、答题时间、答题用时、正确性和得分，为后续诊断与推荐提供依据。",
        "系统不只保存错题结果，还记录错误位置、错误类型、可信程度和复核结论。",
        "推荐结果必须通过后续练习再次验证，避免把一次正确直接判断为已经稳定掌握。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "4 第一版业务模块总览", 1)
    add_module_overview(doc)

    add_heading(doc, "5 主要业务功能", 1)
    add_heading(doc, "5.1 学习档案与学习上下文", 2)
    for text in [
        "维护学生姓名、年级、学期、默认教材版本等学习所需信息。",
        "学校和班级为可选组织关系；未加入组织时仍可正常学习，加入后可继承班级设置。",
        "持续保存练习、作答、错题、掌握度、推荐、作业、闯关和对战记录，形成个人学习档案。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.2 教材、章节与知识点", 2)
    for text in [
        "教材版本用于区分教材名称、出版社、适用年级、学期和版本信息。",
        "一个教材版本包含多个教材章节；教材章节覆盖相应知识点。",
        "同一知识点可以出现在不同教材版本的不同章节中，从而支持通用平台下的多教材适配。",
        "题目通过考查关系关联知识点，系统依据学生所选教材、当前章节和掌握度匹配题目。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.3 题库与出题", 2)
    for text in [
        "维护题目题干、题型、标准答案、解析、难度、知识点、适用年级和使用状态。",
        "支持单选、填空、判断、口算、竖式计算和应用题等小学数学常见题型。",
        "支持教师选题、规则组题和系统个性化组题；组题条件包括知识点、难度、题型、题量和近期练习情况。",
        "同一题目可用于自主练习、教师作业、闯关和对战，但不同场景保留独立的作答与得分记录。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.4 个性化口算与常规练习", 2)
    for text in [
        "个性化口算根据年级、近期错题、答题速度和知识点掌握度生成不同题型、难度和数量的练习。",
        "学生也可主动选择知识点、题型、难度、题量和时间，开展专项练习或限时练习。",
        "练习与题目为多对多包含关系；题目顺序和题目分值属于具体练习中的配置。",
        "每道题的作答结果进入学习记录，并参与错因分析和掌握度更新。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.5 拍照识别与辅助批改", 2)
    for text in [
        "学生拍摄或上传作业图片，系统完成图像处理、题目切分、内容识别和答案提取。",
        "对支持的题型判断正误并生成批改结果；对竖式计算等结构明确的题型，可进一步定位部分步骤错误。",
        "批改结果保存标准答案、学生答案、正确性、可信程度和处理状态。",
        "识别或判断可信程度不足时，提示学生确认；需要教学判断的结果进入教师复核，避免直接写入错误结论。",
        "第一版不承诺识别所有手写应用题的完整推理过程，也不以辅助批改替代教师判断。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.6 错因诊断、掌握度与个性化推荐", 2)
    for text in [
        "按知识点和错误类型归纳错题，记录错误位置、错因说明、可信程度和复核结果。",
        "综合正确率、答题速度、题目难度、练习频率和最近练习时间更新知识点掌握度。",
        "针对薄弱知识点、常见错因和长期未练内容生成推荐计划，安排相似题、变式题和复习题。",
        "学生完成推荐练习后进行再次验证，并据此继续更新掌握度和后续推荐。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "5.7 教师作业与学习报告", 2)
    for text in [
        "教师创建作业并设置标题、题目、分值、适用班级、开始时间和截止时间。",
        "支持统一作业和按学生掌握情况生成的分层作业，记录提交状态、得分、正确率和完成时间。",
        "提供班级和学生两个视角的完成情况、错题分布、知识点掌握情况与阶段趋势。",
        "教师可以根据报告制定训练计划，并在确认后发布给学生。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "6 游戏闯关与学习激励", 1)
    add_heading(doc, "6.1 单人闯关", 2)
    for text in [
        "单人模式统一称为“单人闯关”，不使用“单人PK”表述。",
        "学生主动发起闯关，系统根据年级、所选知识点和当前掌握度生成固定数量题目；第一版默认每关10题，并设置总答题时限。",
        "学生在规定时间内完成全部题目且正确率达到80%，判定闯关成功；题量、时限和通过正确率均可按关卡配置。",
        "闯关成功后记录成绩、增加积分和经验并解锁下一关；失败同样保留记录，展示错题和错因，生成针对性练习后允许再次挑战。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "6.2 游戏化激励", 2)
    for text in [
        "学生完成练习、作业、推荐任务和闯关后获得积分与经验，逐步提升等级。",
        "通过徽章、成就、连续学习天数、每日任务和排行榜展示阶段成果。",
        "奖励规则以鼓励持续学习为主，错误作答不采用强惩罚；系统将错题转化为后续训练机会。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "7 一对一与组队对战", 1)
    add_heading(doc, "7.1 一对一对战", 2)
    for text in [
        "支持系统匹配、同学邀请和房间对战，按照年级、知识点范围和能力水平匹配双方。",
        "双方使用相同题目或同知识点、同难度的等价题目；以正确题数为主要得分依据，用时作为同分判定依据。",
        "保存参赛方、每题作答、实时得分、最终结果、异常状态和结算时间。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "7.2 组队对战", 2)
    for text in [
        "支持队伍创建、加入、退出、成员管理和队伍匹配，第一版以相同人数队伍进行对战。",
        "团队成绩汇总队员答题得分，同时保留每名学生的正确题数、个人得分和贡献排名。",
        "对战结束后结果同时进入个人学习记录，用于发现薄弱知识点并继续推荐针对性练习。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "8 平台运营管理", 1)
    for text in [
        "组织管理：维护学校、年级、班级以及学生入班、离班等组织关系。",
        "内容管理：维护教材版本、章节、知识点、题型、题目、答案、难度和错因分类。",
        "规则管理：维护练习、推荐、闯关、积分、经验、等级、徽章、排行榜和对战规则。",
        "业务处理：查看异常作答、低可信批改、复核记录、对战异常和开放平台使用情况。",
        "运营统计：查看学生规模、学习活跃度、练习完成量、作业完成率、闯关参与度和对战情况。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "9 开放平台", 1)
    add_body(
        doc,
        "开放平台是第一版正式业务模块，不是附加说明。第三方学校系统、教育应用或学习工具可以申请接入，在审核通过的能力范围内调用本平台服务，并查看自身的使用记录和额度情况。",
    )
    add_heading(doc, "9.1 第三方应用管理", 2)
    for text in [
        "提交应用资料、使用场景和所需能力，由管理员完成接入审核、启用、停用和状态管理。",
        "为不同应用配置可使用的业务能力、调用额度和数据范围，避免应用访问与自身业务无关的数据。",
        "记录应用调用时间、调用能力、处理结果和异常信息，便于第三方查看使用情况和定位问题。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "9.2 第一版开放能力", 2)
    for text in [
        "题目查询与组题：按年级、知识点、题型、难度和题量获取题目或生成练习。",
        "答案校验与辅助批改：提交题目和学生答案，返回标准答案、正确性、得分和解释信息。",
        "图片识别与作业切分：提交作业图片，返回识别出的题目、答案和处理状态。",
        "错因诊断与掌握度查询：获取错误分类、知识点掌握情况和学习变化。",
        "个性化推荐与针对性练习：根据学生学习记录获得推荐知识点、推荐题目和训练计划。",
        "智能讲解：在明确题目、答案和知识点的前提下生成适合小学生理解的辅助解释。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "10 第一版主要业务流程", 1)
    flows = [
        "独立学习流程：学生选择年级、学期和教材版本 → 选择或获得练习 → 作答与批改 → 查看错题和掌握度 → 完成推荐练习 → 再次验证。",
        "班级作业流程：教师选择班级和内容范围 → 创建并发布作业 → 学生提交 → 系统统计 → 教师复核与分析 → 发布后续训练。",
        "单人闯关流程：学生选择关卡 → 系统生成题目 → 限时作答 → 判断成功或失败 → 发放奖励或生成巩固练习 → 解锁或再次挑战。",
        "对战流程：学生或队伍发起匹配 → 创建对局与房间 → 公平出题 → 实时作答计分 → 结算 → 写入个人学习记录。",
        "开放调用流程：第三方应用申请接入 → 管理员审核能力范围 → 应用调用业务能力 → 平台返回结果 → 双方查看调用与异常记录。",
    ]
    for text in flows:
        add_list_item(doc, text, decimal_id)

    add_heading(doc, "11 第一版业务边界", 1)
    for text in [
        "系统面向多学校和独立学生，不为某一所学校单独设计业务模型。",
        "学校和班级不是学生使用核心学习功能的前置条件。",
        "家长不作为第一版独立业务角色；第三方开发者以应用主体接入开放平台。",
        "拍照辅助批改优先支持结构明确、可校验的小学数学题型，不承诺覆盖所有手写推理过程。",
        "面向学生的智能能力用于解释、诊断和推荐，不提供无限制开放聊天，也不自动代写作业。",
        "教师发布的作业、训练计划和需要教学判断的批改结果，应保留教师确认环节。",
        "本文档只规定第一版业务功能，不描述具体技术架构、认证方案、数据库表或部署方式。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "12 第一版完成标准", 1)
    for text in [
        "学生能够在不加入学校的情况下完成从练习到再次验证的完整学习闭环。",
        "教师能够完成班级管理、作业发布、结果查看、低可信结果复核和训练计划发布。",
        "管理员能够维护核心内容、业务规则、组织数据和开放平台应用。",
        "单人闯关、一对一对战和组队对战均能形成完整记录、计分与结算结果。",
        "第三方应用能够在审核后的能力范围内调用至少一组学习业务能力，并查询使用结果。",
        "各学习场景产生的作答记录能够汇总到学生学习档案，支持错因、掌握度和推荐持续更新。",
    ]:
        add_list_item(doc, text, bullet_id)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
